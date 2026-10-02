# -*- coding: utf-8 -*-
"""Paramètres, journaux, plan de comptes et règles d'affectation — TOUT est modifiable
dans le classeur regles/REGLES.xlsx (feuilles JOURNAUX, FORMAT, PARAMETRES, REGLES…)."""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill
from .outils import norm

VERSION_REGLES = "2026.10.2"

# ------------------------------------------------------------------ valeurs par défaut
PARAMETRES = {
    "version_regles": VERSION_REGLES,
    "dossier": "MON DOSSIER",
    # comptes
    "compte_fournisseur": "40110000", "tiers_fournisseur": "401199",
    "compte_client": "41110000", "tiers_client": "411101",
    "compte_caisse": "57110000", "compte_banque": "521100",
    "compte_ventes": "70200000", "compte_ventes_2": "70100000",
    "compte_brut": "66100000", "compte_transport": "66340000",
    "compte_cnps_patronal": "66410000", "compte_its_employeur": "64130000",
    "compte_fdfp_charge": "64140000", "compte_personnel": "42200000",
    "compte_avances": "42110000", "compte_cnps": "43100000",
    "compte_its": "44720000", "compte_fdfp": "44780000",
    "compte_timbre": "44730000", "compte_autres_taxes": "64620000",
    "compte_frais_bancaires": "63100000",
    "compte_produits_divers": "75800000", "compte_charges_diverses": "65800000",
    "compte_tva_collectee": "44310000", "compte_tva_deductible": "44510000",
    "compte_attente": "47180000",     # compte d'attente pour les lignes non affectées
    "compte_proprietaire": "46210000",
    # comportement
    "mode_achats": "miroir",          # miroir = 060 charge/401 + règlement ; direct = une seule écriture
    "mode_banque": "direct",
    "ventilation_tva": "non",
    "seuil_immobilisation": "250000",
    "reglement_paie": "caisse",       # caisse | banque | aucun
    "regle_tresorerie": "HORS",       # HORS | OD (en compte propriétaire)
    "recettes_sans_facture": "ATTENTE",  # ATTENTE | HORS | PRODUITS
    "date_min": "", "date_max": "",   # période à traiter (aaaa-mm-jj)
    "modele_comptes": "8",            # 8 chiffres (SYSCOHADA FKS) | 6 chiffres (PCG court)
    # paie (barème 2024 — modifiable)
    "taux_cnps_salarie": "6.3", "cmu_salarie": "500", "taux_cnps_retraite": "7.7",
    "plafond_cnps": "70000", "taux_pf": "5", "taux_maternite": "0.75", "taux_at": "3",
    "cmu_patronale": "500", "taux_its": "16", "abattement_its": "12000",
    "ricf_par_part": "5500", "taux_its_employeur": "1.2", "taux_fdfp": "1.6",
}

# ------------------------------------------------------- les journaux (ordre = ordre des fichiers)
DEFAUT_JOURNAUX = [
    # clé, actif, code journal, nom du journal, préfixe des n° de pièces, nom du fichier
    ("ventes", "oui", "070", "VENTES", "VE", "{n}-A-IMPORTER-Ventes-{code}.txt"),
    ("achats", "oui", "060", "ACHATS", "AC", "{n}-A-IMPORTER-Achats-{code}.txt"),
    ("caisse", "oui", "057", "CAISSE", "CA", "{n}-A-IMPORTER-Caisse-{code}.txt"),
    ("banque", "oui", "052", "BANQUE", "BQ", "{n}-A-IMPORTER-Banque-{code}.txt"),
    ("paie", "oui", "042", "PAIE", "PA", "{n}-A-IMPORTER-Paie-{code}.txt"),
    ("od", "oui", "090", "OPERATIONS DIVERSES", "OD", "{n}-A-IMPORTER-OD-{code}.txt"),
]

# ------------------------------------------------------- le format d'import (A à Z)
DEFAUT_FORMAT = [
    # paramètre, valeur, valeurs possibles
    ("separateur", "TABULATION", "TABULATION · POINT_VIRGULE · VIRGULE · ESPACE · BARRE"),
    ("fin_ligne", "RETOUR_CHARIOT", "RETOUR_CHARIOT (CR) · SAUT_LIGNE (CR+LF) · LF"),
    ("encodage", "cp1252", "cp1252 (Windows) · utf-8 · latin-1"),
    ("format_date", "jj/mm/aaaa", "jj/mm/aaaa · aaaa-mm-jj · jjmmaaaa · jj-mm-aaaa · jj.mm.aaaa"),
    ("decimales", "2", "0 · 2 · 3"),
    ("sep_decimal", "VIRGULE", "VIRGULE · POINT"),
    ("colonnes", "journal|date|piece|compte|tiers|libelle|debit|credit",
     "ordre et choix des colonnes (séparées par | ; ou ,) : journal, date, piece, compte, tiers, libelle, debit, credit"),
    ("entete", "non", "oui = première ligne d'intitulés, non = aucune"),
    ("longueur_libelle", "60", "nombre de caractères maximum du libellé"),
    ("nom_fichier", "{n}-A-IMPORTER-{nom}-{code}.txt",
     "{n} numéro d'ordre · {nom} nom du journal · {code} code journal · {dossier}"),
    ("creer_essais", "oui", "oui = un petit fichier d'essai par journal dans ESSAIS/"),
    ("creer_version_crlf", "oui", "oui = une copie CRLF dans AUTRES/version_CRLF/"),
    ("creer_classeur", "oui", "oui = un classeur Excel de lecture dans AUTRES/"),
    ("creer_rapport", "oui", "oui = le rapport de contrôles dans AUTRES/"),
    ("prefixe_fichier", "", "texte ajouté devant chaque nom de fichier (ex. « FKS_ »)"),
]

# ordre de lecture : la première règle qui colle gagne
DEFAUT_REGLES = [
    # mots-clés, compte (ou [HORS] [IMPOSITION] [PAIE] [IMMOBILISATION] [TRESORERIE]), exclusions
    ("PARALLELISME|CONTRAVENTION|INFRACTION|AMENDE", "[HORS]", ""),
    ("PRET|DON |SOUTIEN|CONDOLEANCES|SOINS|INSPECTION|MISSION|AIDE ", "[HORS]", ""),
    ("DJEBOTEAU|ENVOI A M|ENVOI POUR LE PROPRI|RETRAIT PROPRI|ENVOI PROPRI", "[PROPRIETAIRE]", ""),
    ("AVANCE SUR SALAIRE|AVANCE SALAIRE|AVANCE SUR PAIE|AVANCE SUR LE SALAIRE", "[AVANCE]", ""),
    ("VERSEMENT|DEPOT |DEPÔT|RETRAIT|REMIS A|TRANSMIS|FOND DE CAISSE|FONDS DE CAISSE|ALIMENTATION CAISSE|"
     "TRANSFERT ORANGE|TRANSFERT WAVE|VIREMENT|ENVOI FONDS", "[TRESORERIE]", ""),
    ("REPARAT|REPRATION|DEPANNAGE|DEPANAGE|VIDANGE|LAVAGE|PANNE|MOTEUR|COURROIE|BATTERIE|FEU |FEUX|"
     "CLIGNOTANT|ALTERNATEUR|ESSUIE|SOUDURE|POUBELLE|REVISION|ENTRETIEN|PIECE DE RECHANGE|"
     "VISITE TECHNIQUE|JANTE|PARE BRISE|AMORTISSEUR|SERRURE|CERURE|CYLINDRE|REMORQUAGE|POULIE|"
     "VENTILATEUR|DIAGNOSTIC|COLLAGE|TURBO|FREIN|EMBRAYAGE|SUSPENSION|CARROSSERIE", "62400000", ""),
    ("TRANSPORT|PEAGE|PARKING|STATIONNEMENT|EXPEDITION|COLIS|ENVOIS|YANGO|DEPLACEMENT|CHARGEMENT|LIVRAISON|"
     "REMISE DE FONDS|ENVOI |VOYAGE|DELEGATION", "61830000", ""),
    ("PATENTE|COTISATION|ENREGISTREMENT|FORMULAIRE DE CONTRAT|PROCURATION|PLASTIFICATION|VIGNETTE|DGI|"
     "IMPOT|TAXE|TIMBRE|ATC|GS1", "[IMPOSITION]", "TRANSPORT|PEAGE|DEPLACEMENT"),
    ("JOURNALIER|PAIE DES|MAIN D|TECHNICIEN|AFFUTAGE|MACON|MECANO|MECANICIEN", "60570000",
     "MACONNERIE|AMENAGEMENT|MENUISERIE|COMPTOIR|ETAGERE|INSTALLATION|PEINTURE"),
    ("SALAIRE|PRIME DE VENTE|GRATIFICATION", "[PAIE]", ""),
    ("PNEU|BAGUETTE DE SOUDURE", "60580000", ""),
    ("POLICE", "61830000", ""),
    ("ASSURANCE|ASSSURANCE", "62500000", ""),
    ("LOYER|LOCATION", "62200000", ""),
    ("GAZ|B50|B38|B12", "60420000", ""),
    ("CARBURANT|ESSENCE|GASOIL|HUILE DE FREIN|HUILE MOTEUR", "60530000", ""),
    ("AMENAGEMENT|INSTALLATION|MACONNERIE|MENUISERIE|CONSTRUCTION|COMPTOIR|ETAGERE|PEINTURE|TUYAU|TOLE|"
     "MOULIN|MACHINE|MATERIEL|TABLE|TASSE|USTENSILES|PLAQUE|THERMOS|REFRIGERATEUR|GPS|CAMERA|MOTO|BALANCE|"
     "MEULES|BARRIQUE|ROULETTES|MULTIPRISE|SCEAUX|CACHE NEZ|BACHE|CLIM|ROBINET|SOUS COMPTEUR|ACCESSOIRES",
     "60580000", ""),
    ("CAFE|ACOMPTE|MATIERE PREMIERE", "60210000", "EMBALLAGE|SACHET EXPRESSO|CARNET DE RECU"),
    ("EMBALLAGE|CAPSULE|SACHET|CARTON|MAQUETTE|BOBINE|FIL DE COUTURE| SAC", "60810000", ""),
    ("NOURRITURE|SUCRE|PACKS D'EAU|PACK D'EAU|EAU EN SACHET|SPAGETTI|RIZ|PAIN|ATTIEKE|POISSON|VIANDE|"
     "ROGNON|BLONDS|COLLATION|LAIT EN POUDRE|RESTAURATION|GINGEMBRE|INFUSETTE|EPICE|ARTICLES ET PRODUITS",
     "63830000", ""),
    ("FACTURE CIE|ELECTRICITE|RECHARGE COMPTEUR|CONSOMMATION ENERGIE", "60520000", ""),
    ("SODECI", "60510000", ""),
    ("FOURNITURE|BUREAU|CARNET|CACHET|IMPRESSION|STICKER|SOURIS|SCOTCH|PHOTOCOPIE|CLE OFFICE|LOGICIEL|"
     "LICENCE|ENVELOPPE", "60460000", ""),
    ("EPI |EXTINCTEUR|BALAIS|TRIANGLE DE SOL", "60430000", ""),
    ("COMMUNICATION|COMMUNCATION|COMMIUNICATION|COMMUNICAITON|WIFI", "61830000", ""),
    ("AFFICHE|AFFICHAGE", "62700000", ""),
    ("COMMISSION", "63220000", ""),
    ("WAVE|FRAIS DE TRANSFERT|FRAIS TRANSFERT|FRAIS DE RETRAIT|FRAIS BANCAIRE|COMMISSION BANCAIRE|AGIOS",
     "63100000", ""),
    ("APPROVISIONNEMENT|KIOSQUE", "60100000", ""),
    ("ACHAT", "60580000", ""),
]

DEFAUT_CATEGORIES = [
    ("STOCK|MATIERE PREMIERE", "60210000"), ("GAZ", "60420000"), ("CARBURANT", "60530000"),
    ("EMBALLAGE", "60810000"), ("ELECTRICITE", "60520000"), ("FOURNITURE DE BUREAU", "60460000"),
    ("LOYER", "62200000"), ("NOURRITURE PERSONNEL", "63830000"),
    ("TRANSPORT|EXPEDITION|COMMUNICATION", "61830000"),
    ("MAIN D'OEUVRE|SALAIRE", "60570000"), ("AMENAGEMENT|MATERIEL|MOBILIER", "60580000"),
    ("COMMISSION", "63220000"), ("FRAIS DE TRANSFERT|FRAIS TRANSFERT", "63100000"),
]

DEFAUT_TAXES = [
    ("ARRANGEMENT", "65820000"), ("ENREGISTREMENT|BAIL|FORMULAIRE|PLASTIFICATION", "64600000"),
    ("TIMBRE", "64620000"), ("TAXE PUBLICITAIRE|PUBLICITE", "64680000"), ("PROCURATION", "63250000"),
    ("ATC|GS1", "64120000"), ("PATENTE|COTISATION|DGI|IMPOT|TAXE|VIGNETTE", "64120000"),
]

DEFAUT_IMMOS = [
    ("MOTOS|SOLDE ACHAT MOTO", "24510000"), ("MOTEUR|VEHICULE|CAMIONNETTE", "24510000"),
    ("INFORMATIQUE|ORDINATEUR|IMPRIMANTE", "24410000"),
    ("AMENAGEMENT|MACONNERIE|MENUISERIE|CONSTRUCTION|BATIMENT", "23800000"),
]

DEFAUT_BANQUE = [
    ("LONACI|CLIENT|VERSE|REMISE|VIREMENT RECU|RECU DE|CHEQUE", "41110000", "RECETTE"),
    ("FRAIS|COMMISSION|AGIOS|TENUE DE COMPTE|COTISATION", "63100000", "CHARGE"),
    ("DGI|IMPOT|ITS|FISC", "44710000", "CHARGE"),
    ("CNPS|CNAPS", "43100000", "CHARGE"),
    ("SALAIRE|PAIE", "42200000", "CHARGE"),
    ("ELECTRICITE|CIE", "40110300", "CHARGE"),
    ("RETRAIT|ESPECES|VIREMENT EMIS|FOURNISSEUR|PAIEMENT", "40110000", "CHARGE"),
]

DEFAUT_TIERS = [
    ("401101", "YAKRO", "40110000", "", "FOURNISSEUR"),
    ("401103", "CIE", "40110000", "", "FOURNISSEUR"),
    ("401199", "FOURNISSEURS DIVERS", "40110000", "", "FOURNISSEUR"),
    ("411101", "CAD COTE D'IVOIRE", "41110000", "2601174G", "CLIENT"),
    ("411199", "CLIENTS DIVERS", "41110000", "", "CLIENT"),
]

FEUILLES = ["PARAMETRES", "JOURNAUX", "FORMAT", "REGLES", "CATEGORIES", "TAXES", "IMMOS", "BANQUE", "TIERS"]
CLE_JOURNAUX = [j[0] for j in DEFAUT_JOURNAUX]


class Regles:
    """Charge (et crée si besoin) le classeur de règles, puis classe les libellés."""

    def __init__(self, chemin):
        self.chemin = chemin
        if not os.path.exists(chemin):
            self.creer()
        self.tiers_ajoutes = []
        self.lire()

    # ------------------------------------------------------------- création du classeur
    def creer(self):
        os.makedirs(os.path.dirname(self.chemin), exist_ok=True)
        wb = openpyxl.Workbook()
        t = Font(bold=True, color="FFFFFF")
        r = PatternFill("solid", fgColor="1F3864")
        wb.remove(wb.active)
        for feuille in FEUILLES:
            if feuille == "PARAMETRES":
                entetes, lignes = ["Clé", "Valeur", "Explication"], [[k, v, ""] for k, v in PARAMETRES.items()]
            elif feuille == "JOURNAUX":
                entetes, lignes = (["Clé", "Actif (oui/non)", "Code journal", "Nom du journal",
                                    "Préfixe des pièces", "Nom du fichier exporté"], [list(j) for j in DEFAUT_JOURNAUX])
            elif feuille == "FORMAT":
                entetes, lignes = ["Paramètre", "Valeur", "Valeurs possibles"], [list(f) for f in DEFAUT_FORMAT]
            elif feuille == "REGLES":
                entetes, lignes = (["Mots-clés (séparés par |)", "Compte 8 chiffres", "À exclure si contient"],
                                   DEFAUT_REGLES)
            elif feuille == "CATEGORIES":
                entetes, lignes = ["Votre colonne COMPTE contient", "Compte"], DEFAUT_CATEGORIES
            elif feuille == "TAXES":
                entetes, lignes = ["Mots-clés", "Compte"], DEFAUT_TAXES
            elif feuille == "IMMOS":
                entetes, lignes = ["Mots-clés", "Compte classe 2"], DEFAUT_IMMOS
            elif feuille == "BANQUE":
                entetes, lignes = ["Mots-clés", "Compte", "Nature"], DEFAUT_BANQUE
            else:
                entetes, lignes = ["Compte tiers", "Nom", "Compte général", "NCC", "Type"], DEFAUT_TIERS
            ws = wb.create_sheet(feuille)
            ws.append(entetes)
            for c in ws[1]:
                c.font = t
                c.fill = r
            for l in lignes:
                ws.append(list(l))
            for i, larg in enumerate((52, 18, 26, 20, 16, 42), 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = larg
        wb.save(self.chemin)

    # ------------------------------------------------------------------------- lecture
    def lire(self):
        wb = openpyxl.load_workbook(self.chemin, data_only=True)
        self.p = dict(PARAMETRES)
        for l in wb["PARAMETRES"].iter_rows(min_row=2, values_only=True):
            if l and l[0] and l[1] is not None:
                self.p[str(l[0]).strip()] = str(l[1]).strip()
        # les journaux (l'ordre des lignes = l'ordre des fichiers produits)
        self.journaux = {}
        if "JOURNAUX" in wb.sheetnames:
            for l in self._liste(wb["JOURNAUX"]):
                cle = norm(str(l[0])).lower()
                if not cle:
                    continue
                self.journaux[cle] = dict(
                    cle=cle, actif=str(l[1] if len(l) > 1 and l[1] is not None else "oui").strip().lower(),
                    code=str(l[2] if len(l) > 2 and l[2] is not None else "").strip(),
                    nom=str(l[3] if len(l) > 3 and l[3] is not None else cle).strip().upper(),
                    prefixe=str(l[4] if len(l) > 4 and l[4] is not None else "").strip().upper(),
                    fichier=str(l[5] if len(l) > 5 and l[5] is not None else "").strip())
        for j in DEFAUT_JOURNAUX:                     # on complète si la feuille est incomplète
            if j[0].lower() not in self.journaux:
                self.journaux[j[0].lower()] = dict(cle=j[0].lower(), actif=j[1], code=j[2], nom=j[3],
                                                   prefixe=j[4], fichier=j[5])
        # le format
        self.format = {}
        if "FORMAT" in wb.sheetnames:
            for l in self._liste(wb["FORMAT"]):
                if l and l[0]:
                    self.format[str(l[0]).strip()] = str(l[1] if len(l) > 1 and l[1] is not None else "").strip()
        for f in DEFAUT_FORMAT:
            self.format.setdefault(f[0], f[1])
        # compatibilité : les paramètres historiques suivent la feuille JOURNAUX
        for cle, j in self.journaux.items():
            self.p["code_" + str(cle).lower()] = j["code"]
            self.p["piece_" + str(cle).lower()] = j["prefixe"]
        self.regles = [(self._mots(a), str(b).strip(), self._mots(c)) for a, b, c in self._liste(wb["REGLES"])]
        self.categories = [(self._mots(a), str(b)) for a, b in self._liste(wb["CATEGORIES"]) if b]
        self.taxes = [(self._mots(a), str(b)) for a, b in self._liste(wb["TAXES"]) if b]
        self.immos = [(self._mots(a), str(b)) for a, b in self._liste(wb["IMMOS"]) if b]
        self.banque = [(self._mots(a), str(b), norm(c)) for a, b, c in self._liste(wb["BANQUE"])]
        self.tiers = []
        for l in self._liste(wb["TIERS"]):
            if l[0]:
                self.tiers.append({"compte": str(l[0]).strip(), "nom": str(l[1] or "").strip(),
                                   "general": str(l[2] or "40110000").strip(),
                                   "ncc": str(l[3] or "").strip(), "type": str(l[4] or "").strip()})

    @staticmethod
    def _liste(ws):
        out = []
        for l in ws.iter_rows(min_row=2, values_only=True):
            if l and any(x is not None and str(x).strip() for x in l):
                out.append([x if x is not None else "" for x in l])
        return out

    @staticmethod
    def _mots(texte):
        """'PNEU|BAGUETTE DE SOUDURE' -> ['PNEU', 'BAGUETTE DE SOUDURE'] normalisés."""
        return [norm(x) for x in str(texte or "").split("|") if x.strip()]

    # ------------------------------------------------------------------- paramètres
    def val(self, cle, defaut=""):
        return self.p.get(cle, defaut)

    def nombre(self, cle):
        from .outils import nombre as _n
        return _n(self.p.get(cle, 0))

    def oui(self, cle, defaut="oui"):
        return str(self.p.get(cle, defaut)).strip().lower() in ("oui", "yes", "o", "1", "vrai", "true")

    # --------------------------------------------------------------------- journaux
    def journal(self, cle):
        return self.journaux.get(str(cle).lower(),
                                 dict(cle=str(cle).lower(), actif="non", code="", nom=str(cle).upper(),
                                      prefixe="", fichier=""))

    def code(self, cle):
        j = self.journal(cle)
        return j["code"] or self.p.get("code_" + str(cle).lower(), "")

    def actif(self, cle):
        j = self.journal(cle)
        return j["actif"] not in ("non", "no", "n", "0", "faux", "false", "")

    def nom_journal(self, cle):
        return self.journal(cle)["nom"]

    def prefixe(self, cle):
        return self.journal(cle)["prefixe"]

    def journaux_actifs(self):
        """Les journaux activés, dans l'ordre de la feuille JOURNAUX (clé libre ou connue)."""
        return [j for j in self.journaux.values() if self.actif(j["cle"])]

    # --------------------------------------------------------------------- format
    def fmt(self, cle, defaut=""):
        return self.format.get(cle, defaut)

    # ------------------------------------------------------------------- affectations
    def classer(self, libelle, categorie=""):
        """Rend (compte ou marqueur, motif). Marqueurs : [HORS] [TRESORERIE] [PAIE] [IMPOSITION]."""
        k = norm(libelle)
        c = norm(categorie)
        for mots, compte, sauf in self.regles:
            if not mots:
                continue
            if any(w in k for w in mots) and not any(w in k for w in sauf):
                return compte, "libellé"
        if c:
            for mots, compte in self.categories:
                if any(w in c for w in mots):
                    return compte, "colonne COMPTE"
        return "", "à préciser"

    def taxe(self, libelle):
        k = norm(libelle)
        for mots, compte in self.taxes:
            if any(w in k for w in mots):
                return compte
        return self.val("compte_charges_diverses", "65800000")

    def immo(self, libelle, montant):
        if montant <= self.nombre("seuil_immobilisation"):
            return ""
        k = norm(libelle)
        for mots, compte in self.immos:
            if any(w in k for w in mots):
                return compte
        return ""

    def banque_nature(self, libelle):
        k = norm(libelle)
        for mots, compte, nature in self.banque:
            if any(w in k for w in mots):
                return compte, nature
        return "", ""

    # ------------------------------------------------------------------- comptes tiers
    def tiers_pour_ncc(self, ncc, nom=""):
        """Retrouve (ou crée) le compte tiers d'un client à partir de son NCC de facture FNE."""
        ncc = str(ncc or "").strip().upper()
        for t in self.tiers:
            if ncc and t["ncc"].upper() == ncc:
                return t
        n = norm(nom)
        for t in self.tiers:
            if n and norm(t["nom"]) and norm(t["nom"]) in n:
                return t
        pris = {t["compte"] for t in self.tiers}
        numero = 411101
        while str(numero) in pris:
            numero += 1
        t = {"compte": str(numero), "nom": (nom or ncc or "CLIENT").upper()[:40],
             "general": self.val("compte_client", "41110000"), "ncc": ncc, "type": "CLIENT"}
        self.tiers.append(t)
        self.tiers_ajoutes.append(t)
        return t

    def tiers_fournisseur(self, nom=""):
        n = norm(nom)
        for t in self.tiers:
            if t["type"].upper() == "FOURNISSEUR" and n and norm(t["nom"]) and norm(t["nom"]) in n:
                return t
        for t in self.tiers:
            if t["compte"] == self.val("tiers_fournisseur", "401199"):
                return t
        return {"compte": self.val("tiers_fournisseur", "401199"), "nom": "FOURNISSEURS DIVERS",
                "general": self.val("compte_fournisseur", "40110000"), "ncc": "", "type": "FOURNISSEUR"}

    def sauver_tiers(self):
        """Écrit dans le classeur de règles les comptes clients créés automatiquement."""
        if not self.tiers_ajoutes:
            return []
        wb = openpyxl.load_workbook(self.chemin)
        ws = wb["TIERS"]
        for t in self.tiers_ajoutes:
            ws.append([t["compte"], t["nom"], t["general"], t["ncc"], t["type"]])
        wb.save(self.chemin)
        return self.tiers_ajoutes

    # ------------------------------------------------- écriture des réglages (interface)
    def sauver_reglages(self, journaux=None, format_=None, parametres=None):
        """Enregistre dans le classeur ce qui a été modifié à l'écran (journaux, format, paramètres)."""
        wb = openpyxl.load_workbook(self.chemin)
        if journaux:
            ws = wb["JOURNAUX"]
            for l in list(ws.iter_rows(min_row=2)):
                cle = norm(str(l[0].value or "")).lower()
                if cle in journaux:
                    d = journaux[cle]
                    for i, champ in enumerate(("actif", "code", "nom", "prefixe", "fichier"), 1):
                        if champ in d and d[champ] is not None:
                            l[i].value = d[champ]
            connues = {norm(str(l[0].value or "")).lower() for l in ws.iter_rows(min_row=2)}
            for cle, d in journaux.items():
                if cle not in connues:                # journal ajouté par l'utilisateur
                    ws.append([d.get("cle", cle).lower(), d.get("actif", "oui"), d.get("code", ""),
                               d.get("nom", ""), d.get("prefixe", ""),
                               d.get("fichier", "{n}-A-IMPORTER-{nom}-{code}.txt")])
        if format_:
            ws = wb["FORMAT"]
            for l in ws.iter_rows(min_row=2):
                cle = str(l[0].value or "").strip()
                if cle in format_:
                    l[1].value = format_[cle]
            connues = {str(l[0].value or "").strip() for l in ws.iter_rows(min_row=2)}
            for cle, v in format_.items():
                if cle not in connues:
                    ws.append([cle, v, ""])
        if parametres:
            ws = wb["PARAMETRES"]
            for l in ws.iter_rows(min_row=2):
                cle = str(l[0].value or "").strip()
                if cle in parametres:
                    l[1].value = parametres[cle]
            connues = {str(l[0].value or "").strip() for l in ws.iter_rows(min_row=2)}
            for cle, v in parametres.items():
                if cle not in connues:
                    ws.append([cle, v, ""])
        wb.save(self.chemin)
        self.lire()
        return True

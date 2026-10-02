# -*- coding: utf-8 -*-
"""Écriture des fichiers d'import Sage, du classeur de lecture et des notes.

TOUT le format vient de la feuille FORMAT du classeur de règles : séparateur, fin de ligne,
encodage, format de date, décimales, ordre et choix des colonnes, en-tête, nom des fichiers.
Les journaux (feuille JOURNAUX) peuvent être activés / désactivés et renommés.
"""
import collections
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .outils import f, jjmm, texte_libre

CHAMPS = ("journal", "date", "piece", "compte", "tiers", "libelle", "debit", "credit")
ENTETES = {"journal": "Code journal", "date": "Date de pièce", "piece": "N° de pièce",
           "compte": "N° compte général", "tiers": "N° compte tiers", "libelle": "Libellé",
           "debit": "Débit", "credit": "Crédit"}
SEPARATEURS = {"TABULATION": "\t", "POINT_VIRGULE": ";", "VIRGULE": ",", "ESPACE": " ", "BARRE": "|"}
FINS = {"RETOUR_CHARIOT": "\r", "SAUT_LIGNE": "\r\n", "LF": "\n"}


# ---------------------------------------------------------------------------- format
def _date(iso, regles):
    if not iso:
        return ""
    j, m, a = iso[8:10], iso[5:7], iso[0:4]
    g = regles.fmt("format_date", "jj/mm/aaaa").strip().lower()
    if g.startswith("aaaa"):
        return "%s-%s-%s" % (a, m, j)
    if g in ("jjmmaaaa", "jj/mm/aaaa sans séparateur", "sans separateur"):
        return "%s%s%s" % (j, m, a)
    if g.startswith("jj-"):
        return "%s-%s-%s" % (j, m, a)
    if g.startswith("jj."):
        return "%s.%s.%s" % (j, m, a)
    return "%s/%s/%s" % (j, m, a)


def _montant(x, regles):
    x = float(x or 0)
    dec = {"0": 0, "2": 2, "3": 3}.get(str(regles.fmt("decimales", "2")).strip(), 2)
    point = regles.fmt("sep_decimal", "VIRGULE").strip().upper().startswith("POINT")
    t = ("%.*f" % (dec, x))
    return t if point else t.replace(".", ",")


def valeur(champ, p, l, regles):
    if champ == "journal":
        return p["journal"]
    if champ == "date":
        return _date(p["date"], regles)
    if champ == "piece":
        return p["numero"] or ""
    if champ == "compte":
        return l["compte"]
    if champ == "tiers":
        return l["tiers"] or ""
    if champ == "libelle":
        return texte_libre(l["libelle"], int(regles.fmt("longueur_libelle", "60") or 60))
    return _montant(l["debit"] if champ == "debit" else l["credit"], regles)


def colonnes(regles):
    brut = str(regles.fmt("colonnes", "|".join(CHAMPS))).replace(";", "|").replace(",", "|")
    pris = [c.strip().lower() for c in brut.split("|")]
    out = [c for c in pris if c in CHAMPS]
    return out or list(CHAMPS)


def ligne(p, l, regles, fin=None):
    """Une ligne du fichier d'import, prête à écrire (octets)."""
    sep = SEPARATEURS.get(regles.fmt("separateur", "TABULATION").strip().upper(), "\t")
    fin = fin or FINS.get(regles.fmt("fin_ligne", "RETOUR_CHARIOT").strip().upper(), "\r")
    enc = regles.fmt("encodage", "cp1252") or "cp1252"
    t = sep.join(valeur(c, p, l, regles) for c in colonnes(regles)) + fin
    return t.encode(enc, errors="replace")


def ligne_entete(regles):
    sep = SEPARATEURS.get(regles.fmt("separateur", "TABULATION").strip().upper(), "\t")
    fin = FINS.get(regles.fmt("fin_ligne", "RETOUR_CHARIOT").strip().upper(), "\r")
    enc = regles.fmt("encodage", "cp1252") or "cp1252"
    t = sep.join(ENTETES[c] for c in colonnes(regles)) + fin
    return t.encode(enc, errors="replace")


def nom_fichier(cle, numero, regles):
    j = regles.journal(cle)
    gabarit = j["fichier"] or regles.fmt("nom_fichier", "{n}-A-IMPORTER-{nom}-{code}.txt")
    nom = re.sub(r"\s+", "-", j["nom"] or cle).strip("-")
    t = (gabarit.replace("{n}", str(numero)).replace("{nom}", nom)
         .replace("{code}", j["code"]).replace("{cle}", cle)
         .replace("{dossier}", re.sub(r"\s+", "-", regles.val("dossier", "DOSSIER"))))
    t = regles.fmt("prefixe_fichier", "") + t
    return re.sub(r"[\\/:*?\"<>|]+", "-", t)


def ecrire_lot(chemin, pieces, regles):
    with open(chemin, "wb") as fh:
        if regles.fmt("entete", "non").strip().lower() in ("oui", "yes", "o", "1"):
            fh.write(ligne_entete(regles))
        for p in pieces:
            for l in p["lignes"]:
                fh.write(ligne(p, l, regles))
    return chemin


def ecrire_tout(sortie, res, regles, blocage=None):
    """Écrit les journaux actifs, les essais, la version CRLF, le classeur et les notes."""
    os.makedirs(sortie, exist_ok=True)
    for sous in ("AUTRES", "ESSAIS"):
        os.makedirs(os.path.join(sortie, sous), exist_ok=True)
    if regles.oui("creer_version_crlf", "oui"):
        os.makedirs(os.path.join(sortie, "AUTRES", "version_CRLF"), exist_ok=True)
    journaux = res["journaux"]
    fichiers, n = [], 0
    for j in regles.journaux_actifs():
        pieces = journaux.get(j["code"])
        if not pieces:
            continue
        n += 1
        nom = nom_fichier(j["cle"], n, regles)
        chemin = os.path.join(sortie, nom)
        ecrire_lot(chemin, pieces, regles)
        fichiers.append((nom, chemin, j["code"], len(pieces),
                         sum(len(p["lignes"]) for p in pieces),
                         sum(l["debit"] for p in pieces for l in p["lignes"])))
        # version avec saut de ligne Windows (si Sage refuse la première)
        if regles.oui("creer_version_crlf", "oui"):
            chemin_crlf = os.path.join(sortie, "AUTRES", "version_CRLF", nom)
            with open(chemin_crlf, "wb") as fh:
                if regles.fmt("entete", "non").strip().lower() in ("oui", "yes", "o", "1"):
                    fh.write(ligne_entete(regles))
                for p in pieces:
                    for l in p["lignes"]:
                        fh.write(ligne(p, l, regles, fin="\r\n"))
        # fichier d'essai (les 3 premières pièces) pour valider le format dans Sage
        if regles.oui("creer_essais", "oui"):
            ecrire_lot(os.path.join(sortie, "ESSAIS", "essai-" + nom), pieces[:3], regles)

    # -------------------------------------------------------------- classeur de lecture
    if regles.oui("creer_classeur", "oui"):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        t, r = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
        ws = wb.create_sheet("SYNTHESE")
        ws.append(["Journal", "Code", "Pièces", "Lignes", "Débit", "Crédit", "Écart",
                   "Pièces déséquilibrées"])
        for c in ws[1]:
            c.font, c.fill = t, r
        noms = {j["code"]: j["nom"] for j in regles.journaux.values()}
        for s in res["stats"]:
            ws.append([noms.get(s["journal"], s["journal"]), s["journal"], s["pieces"], s["lignes"],
                       s["debit"], s["credit"], s["ecart"], s["ko"]])
        for i, larg in enumerate((22, 8, 10, 10, 16, 16, 9, 20), 1):
            ws.column_dimensions[get_column_letter(i)].width = larg
        for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
            for c in row:
                c.number_format = "#,##0"
        for code, pieces in journaux.items():
            nom_feuille = (noms.get(code, code) or code)[:31]
            w = wb.create_sheet(nom_feuille)
            w.append(["Journal", "Date", "N° pièce", "Compte", "Tiers", "Libellé", "Débit", "Crédit"])
            for c in w[1]:
                c.font, c.fill = t, r
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, larg in enumerate((9, 12, 20, 14, 12, 44, 14, 14), 1):
                w.column_dimensions[get_column_letter(i)].width = larg
            for p in pieces:
                for l in p["lignes"]:
                    w.append([p["journal"], _date(p["date"], regles), p["numero"], l["compte"],
                              l["tiers"], l["libelle"], l["debit"], l["credit"]])
            for row in w.iter_rows(min_row=2, min_col=7, max_col=8):
                for c in row:
                    c.number_format = "#,##0"
            w.freeze_panes = "A2"
        wb.save(os.path.join(sortie, "AUTRES", "1-Journaux_a_lire.xlsx"))

    # ------------------------------------------------------------------------ rapport
    if regles.oui("creer_rapport", "oui"):
        with open(os.path.join(sortie, "AUTRES", "RAPPORT_CONTROLES.txt"), "w", encoding="utf-8") as fh:
            fh.write("CONTRÔLES DES JOURNAUX — %s\n" % regles.val("dossier", ""))
            fh.write("=" * 96 + "\n")
            fh.write("%-22s %8s %8s %16s %16s %8s\n" % ("JOURNAL", "PIÈCES", "LIGNES", "DÉBIT", "CRÉDIT",
                                                        "ÉCART"))
            for s in res["stats"]:
                fh.write("%-22s %8d %8d %16s %16s %8.2f%s\n"
                         % ((noms.get(s["journal"], s["journal"]) + " (" + s["journal"] + ")"),
                            s["pieces"], s["lignes"], f(s["debit"]), f(s["credit"]), s["ecart"],
                            "   <-- À CORRIGER" if abs(s["ecart"]) > 0.01 else ""))
            tot_d = sum(s["debit"] for s in res["stats"])
            tot_c = sum(s["credit"] for s in res["stats"])
            fh.write("-" * 96 + "\n%-22s %8d %8d %16s %16s %8.2f\n"
                     % ("TOTAL", sum(s["pieces"] for s in res["stats"]),
                        sum(s["lignes"] for s in res["stats"]), f(tot_d), f(tot_c), tot_d - tot_c))
            fh.write("\n%s\n" % ("TOUT EST ÉQUILIBRÉ" if abs(tot_d - tot_c) < 0.01
                                 else "ATTENTION : il reste un écart à corriger"))
            fh.write("\nJOURNAUX — état demandé dans la feuille JOURNAUX\n" + "-" * 96 + "\n")
            for j in regles.journaux.values():
                fh.write("   %-20s code %-4s %-6s %s\n" % (j["nom"], j["code"],
                                                           "ACTIF" if regles.actif(j["cle"]) else "désactivé",
                                                           j["fichier"]))
            if res["alertes"]:
                fh.write("\nPOINTS À VÉRIFIER\n" + "-" * 96 + "\n")
                for a in res["alertes"]:
                    fh.write("  - %s\n" % a)
            if res["a_confirmer"]:
                fh.write("\nLIGNES À AFFECTER (compte d'attente %s en attendant votre réponse)\n"
                         % regles.val("compte_attente", "47180000") + "-" * 96 + "\n")
                for c in res["a_confirmer"]:
                    fh.write("  %-11s %12s  %-44s %s\n" % (c["quoi"], f(c["montant"]),
                                                           (c["libelle"] or "")[:44], c["cle"]))
            if res["hors"]:
                tot = sum(x["depense"] or x["entree"] for x in res["hors"])
                fh.write("\nHORS PÉRIMÈTRE (%d lignes, %s F) — non journalisées, pour information\n"
                         % (len(res["hors"]), f(tot)) + "-" * 96 + "\n")
                grou = collections.defaultdict(lambda: [0, 0.0])
                for x in res["hors"]:
                    grou[x["motif"]][0] += 1
                    grou[x["motif"]][1] += x["depense"] or x["entree"]
                for motif, (nb, s) in sorted(grou.items(), key=lambda y: -y[1][1]):
                    fh.write("  %-58s %4d lignes  %14s F\n" % (motif[:58], nb, f(s)))
                for x in sorted(res["hors"], key=lambda y: -(y["depense"] or y["entree"]))[:400]:
                    fh.write("      %-10s %14s  %s\n" % (_date(x["date"], regles), f(x["depense"] or x["entree"]),
                                                         (x["libelle"] or "")[:60]))

    # ------------------------------------------------------------------ tiers à créer
    with open(os.path.join(sortie, "COMPTES_TIERS_A_CREER.txt"), "w", encoding="utf-8") as fh:
        fh.write("COMPTES TIERS UTILISÉS PAR LES JOURNAUX — à créer dans Sage s'ils n'existent pas\n")
        fh.write("=" * 96 + "\n")
        fh.write("%-12s %-38s %-14s %-12s %s\n" % ("COMPTE TIERS", "NOM", "COMPTE GÉNÉRAL", "NCC", "TYPE"))
        for x in regles.tiers:
            fh.write("%-12s %-38s %-14s %-12s %s\n" % (x["compte"], x["nom"][:38], x["general"], x["ncc"],
                                                       x["type"]))
        if regles.tiers_ajoutes:
            fh.write("\nCréés automatiquement à partir des factures FNE : %s\n"
                     % ", ".join("%s (%s)" % (x["compte"], x["nom"]) for x in regles.tiers_ajoutes))
    return fichiers


def lisez_moi(sortie, res, regles, nom_dossier="", fichiers=()):
    p = regles.p
    lignes = ["%s — JOURNAUX POUR SAGE 100" % p.get("dossier", "DOSSIER"), "=" * 96, ""]
    for nom, chemin, code, nb, lig, tot in fichiers:
        lignes.append("   %-42s %4d pièces %5d lignes %16s F" % (nom, nb, lig, f(tot)))
    lignes += ["", "Comment importer", "-" * 96,
               "   1. Sage : Échanges > Importer des écritures, avec votre format d'import.",
               "   2. Supprimez les règles 401* et 411* du format si Sage refuse les tiers.",
               "   3. Importez les fichiers dans l'ordre 1, 2, 3… ; ne réimportez jamais un fichier "
               "(doublons).",
               "   4. En cas de refus, essayez la version du dossier AUTRES/version_CRLF.",
               "", "Les journaux de ce lot", "-" * 96]
    noms = {j["code"]: j["nom"] for j in regles.journaux.values()}
    for s in res["stats"]:
        lignes.append("   %s (%s) : %d pièces, %s F au débit = %s F au crédit (écart %.2f)"
                      % (noms.get(s["journal"], s["journal"]), s["journal"], s["pieces"], f(s["debit"]),
                         f(s["credit"]), s["ecart"]))
    for j in regles.journaux.values():
        if not regles.actif(j["cle"]):
            lignes.append("   %s (%s) : DÉSACTIVÉ — aucun fichier produit" % (j["nom"], j["code"]))
    t = res["totaux"]
    if t["factures"]:
        lignes += ["", "Ventes", "-" * 96,
                   "   %d factures FNE lues, %s F de ventes HT, %s F de taxes/timbre."
                   % (t["factures"], f(t["ventes_ht"]), f(t["ventes_taxe"]))]
    if t["bulletins"]:
        lignes += ["", "Paie", "-" * 96,
                   "   %d bulletins lus sur %d mois ; règlement des nets par %s."
                   % (t["bulletins"], t["mois_paie"], regles.val("reglement_paie", "caisse"))]
    lignes += ["", "Points à vérifier", "-" * 96]
    for a in (res["alertes"][:12] or ["   (rien à signaler)"]):
        lignes.append("   - %s" % a)
    if res["a_confirmer"]:
        lignes.append("   - %d ligne(s) à affecter (compte d'attente %s) : voir AUTRES/RAPPORT_CONTROLES.txt"
                      % (len(res["a_confirmer"]), regles.val("compte_attente", "47180000")))
    if res["hors"]:
        lignes.append("   - %d ligne(s) hors périmètre, listées dans AUTRES/RAPPORT_CONTROLES.txt"
                      % len(res["hors"]))
    lignes += ["", "Format d'export (feuille FORMAT du classeur de règles)", "-" * 96,
               "   Séparateur %s · fin de ligne %s · encodage %s · dates %s · %s décimales (%s) · "
               "colonnes : %s · en-tête %s"
               % (regles.fmt("separateur"), regles.fmt("fin_ligne"), regles.fmt("encodage"),
                  regles.fmt("format_date"), regles.fmt("decimales"), regles.fmt("sep_decimal"),
                  regles.fmt("colonnes"), regles.fmt("entete")), ""]
    open(os.path.join(sortie, "LISEZ-MOI.txt"), "w", encoding="utf-8").write("\n".join(lignes))

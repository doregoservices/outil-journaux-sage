# -*- coding: utf-8 -*-
"""Lecture d'un registre de factures FNE exporté en Excel (bouton « Mes factures » du portail DGI).

C'est le même contenu que les PDF FNE, en plus pratique : une ligne par facture.
"""
import os
import openpyxl
from .outils import norm, nombre, date_fr

SYN = {
    "numero": ["REFERENCE", "N FACTURE", "NUMERO", "NO FACTURE", "NUMERO DE FACTURE", "FACTURE"],
    "date": ["DATE"],
    "mode": ["MODE DE PAIEMENT", "MODE PAIEMENT", "REGLEMENT"],
    "timbre": ["TIMBRE DE QUITTANCE", "TIMBRE"],
    "ht": ["TOTAL HT", "MONTANT HT", "TOTAL HORS TAXE"],
    "tva": ["TOTAL TVA", "MONTANT TVA"],
    "autres": ["TOTAL AUTRES TAXES", "AUTRES TAXES"],
    "total": ["NET A PAYER", "NET A PAYER TTC", "TOTAL TTC", "TOTAL A PAYER", "MONTANT TTC"],
    "ncc": ["NCC DU CLIENT", "NCC CLIENT", "NCC"],
    "client": ["NOM DE LA SOCIETE DU CLIENT", "NOM DE LA SOCIETE", "NOM DU CLIENT", "CLIENT",
               "RAISON SOCIALE", "NOM CLIENT"],
    "type": ["TYPE DE FACTURE", "TYPE FACTURE"],
    "soustype": ["SOUS TYPE DE FACTURE", "SOUS TYPE", "STATUT"],
    "vendeur": ["ETABLISSEMENT", "NOM DU VENDEUR", "POINT DE VENTE"],
}
MODES = {"check": "Chèque", "cash": "Espèces", "especes": "Espèces", "espèces": "Espèces",
         "mobile_money": "Mobile Money", "transfer": "Virement", "card": "Carte"}


def _entete(ws):
    """Trouve la ligne d'en-têtes et associe chaque colonne. On privilégie les intitulés
    exacts, puis on complète par ressemblance (« Référence initial » ne doit pas prendre
    la place de « Référence »)."""
    for i, ligne in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True)):
        cellules = [norm(c) for c in ligne]
        cols = {}
        for exact in (True, False):
            for cle, motifs in SYN.items():
                if cle in cols:
                    continue
                for j, c in enumerate(cellules):
                    if not c:
                        continue
                    if any(c == m for m in motifs) if exact else \
                       any(len(m) > 5 and (c.startswith(m) or m in c) for m in motifs):
                        cols[cle] = j
                        break
        if "numero" in cols and ("total" in cols or "ht" in cols) and "date" in cols:
            return i, cols
    return None, {}


def est_un_registre(chemin):
    try:
        wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        for ws in wb.worksheets:
            i, cols = _entete(ws)
            if i is not None:
                wb.close()
                return True
        wb.close()
    except Exception:
        return False
    return False


def lire(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    factures, alertes = [], []
    for ws in wb.worksheets:
        lig, cols = _entete(ws)
        if lig is None:
            continue
        for i, r in enumerate(ws.iter_rows(min_row=lig + 2, values_only=True), lig + 2):
            val = lambda cle: (r[cols[cle]] if cle in cols and cols[cle] < len(r) else None)
            numero = str(val("numero") or "").strip()
            if not numero or norm(numero) in ("REFERENCE", "N FACTURE"):
                continue
            d = date_fr(val("date"))
            ht = nombre(val("ht"))
            tva = nombre(val("tva"))
            autres = nombre(val("autres"))
            timbre = nombre(val("timbre"))
            total = nombre(val("total")) or (ht + tva + autres + timbre)
            if not total:
                continue
            client = str(val("client") or "").strip()
            mode = str(val("mode") or "").strip()
            factures.append(dict(type="facture_vte", fichier=os.path.basename(chemin),
                                 numero=numero, date=d.isoformat() if d else "",
                                 client_nom=client, client_ncc=str(val("ncc") or "").strip().upper(),
                                 mode_paiement=MODES.get(norm(mode).lower(), mode),
                                 articles=[], ht=ht, tva=tva, autres_taxes=autres + timbre,
                                 timbre=timbre, total_a_payer=total, exo=(tva == 0),
                                 statut=str(val("soustype") or ""), feuille=ws.title, ligne=i))
        break
    if not factures:
        alertes.append("Registre de factures reconnu mais aucune ligne lisible dans %s."
                       % os.path.basename(chemin))
    statuts = {norm(f["statut"]) for f in factures if f["statut"]}
    if statuts - {"INVOICE", "FACTURE", ""}:
        alertes.append("Types de factures présents : %s — vérifiez s'il faut écarter les proformas/avoirs."
                       % ", ".join(sorted(statuts)))
    return dict(type="ventes", fichier=os.path.basename(chemin), factures=factures, alertes=alertes)
